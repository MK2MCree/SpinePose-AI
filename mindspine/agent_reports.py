"""Bounded local report exports. Rendering uses recorded case facts, never a model."""
from __future__ import annotations

from datetime import date, datetime
from html import escape
from pathlib import Path
from typing import Any
import importlib.util
import io
import os
import re
import sys
import sysconfig

REPORT_VERSION = "spine-case-export-v1"
SECTIONS = {"profile": "个人资料", "questionnaire": "问卷与日常功能", "assessments": "检测与模型结果", "followup": "随访与复核"}
DEFAULTS = {"title": "脊柱健康评估报告", "sections": list(SECTIONS), "audience": "patient", "detail": "full", "format": "both", "date_from": None, "date_to": None}
_ROOT = Path(__file__).resolve().parents[1]


def _dependencies() -> None:
    for name in ("pdf", "reports"):
        path = _ROOT / ".runtime" / name
        if path.is_dir() and str(path) not in sys.path:
            # lxml is a native dependency: prefer the local install only when
            # its extension matches this interpreter (QA also uses Python 3.12).
            suffix = sysconfig.get_config_var("EXT_SUFFIX") or ""
            compatible = name == "pdf" or bool(suffix and list((path / "lxml").glob("etree*" + suffix)))
            if compatible: sys.path.insert(0, str(path))
            else: sys.path.append(str(path))


def report_capabilities() -> dict[str, Any]:
    _dependencies()
    return {"version": REPORT_VERSION, "formats": [f for f, pkg in (("pdf", "reportlab"), ("docx", "docx")) if importlib.util.find_spec(pkg)],
            "sections": [{"id": key, "title": value} for key, value in SECTIONS.items()], "defaults": dict(DEFAULTS)}


def validate_report_options(options: dict | None) -> dict:
    if options is not None and not isinstance(options, dict):
        raise ValueError("报告选项必须为对象")
    supplied = options or {}
    if set(supplied) - set(DEFAULTS):
        raise ValueError("包含不支持的报告选项")
    result = {**DEFAULTS, **supplied}
    title = result["title"]
    if not isinstance(title, str) or not title.strip() or len(title) > 80 or any(ord(c) < 32 for c in title):
        raise ValueError("标题须为 1 至 80 个可见字符")
    result["title"] = title.strip()
    sections = result["sections"]
    if not isinstance(sections, list) or not sections or len(sections) > 4 or any(not isinstance(s, str) or s not in SECTIONS for s in sections):
        raise ValueError("请选择支持的报告内容范围")
    result["sections"] = list(dict.fromkeys(sections))
    for key, allowed in (("audience", {"patient", "clinician"}), ("detail", {"brief", "full"}), ("format", {"pdf", "docx", "both"})):
        if not isinstance(result[key], str) or result[key] not in allowed:
            raise ValueError("报告选项不受支持：" + key)
    for key in ("date_from", "date_to"):
        value = result[key]
        if value in (None, ""):
            result[key] = None
            continue
        if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
            raise ValueError("日期须为 YYYY-MM-DD")
        try:
            date.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("报告日期无效") from exc
    if result["date_from"] and result["date_to"] and result["date_from"] > result["date_to"]:
        raise ValueError("开始日期不能晚于结束日期")
    return result


def parse_report_request(message: str) -> dict | None:
    """Recognise creation intent, not a request to explain an existing report."""
    text = str(message or "").strip()
    # Evaluate clauses independently: “不用导出 PDF，请生成 Word 报告”
    # must not silently re-enable the explicitly declined PDF format.
    clauses = re.split(r"[，,。；;！？!?\n]", text)
    text = "，".join(clause for clause in clauses if not re.search(
        r"(?:不要|不用|无需|别|暂不|不想|不需要|取消).{0,30}(?:生成|导出|制作|做|下载|报告|档案|pdf|word|docx)|(?:报告|档案|pdf|word|docx).{0,8}(?:不需要|不用|取消|不要)", clause, re.I))
    subject = r"(?:报告|档案|pdf|docx|word)"
    if not re.search(r"(?:生成|导出|制作|下载|做|整理成|出(?:一)?份|给我一份|保存为).{0,25}" + subject, text, re.I):
        return None
    result = dict(DEFAULTS)
    if "档案" in text and "报告" not in text:
        result["title"] = "脊柱健康档案"
    if re.search(r"医生|医师|专业版", text): result["audience"] = "clinician"
    if re.search(r"简版|简短|简洁|摘要", text): result["detail"] = "brief"
    has_pdf = bool(re.search(r"pdf", text, re.I))
    has_docx = bool(re.search(r"docx|word", text, re.I))
    if has_pdf != has_docx: result["format"] = "pdf" if has_pdf else "docx"
    if re.search(r"仅|只", text):
        found = [key for key, pattern in (("profile", "个人资料"), ("questionnaire", "问卷"), ("assessments", "检测|影像|测量"), ("followup", "随访|打卡")) if re.search(pattern, text)]
        if found: result["sections"] = found
    title_match = re.search(r"标题(?:为|是|叫|设为)?[：:\s]*[“\"]([^”\"]{1,80})[”\"]", text)
    if title_match: result["title"] = title_match.group(1)
    return validate_report_options(result)


def _text(value: Any, limit: int = 1800) -> str:
    if value is None or value == "" or value == [] or value == {}: return "未记录"
    if isinstance(value, bool): return "是" if value else "否"
    if isinstance(value, (list, tuple)): return "；".join(_text(v, 300) for v in value[:30])[:limit]
    if isinstance(value, dict): return "；".join(f"{k}：{_text(v, 300)}" for k, v in list(value.items())[:24])[:limit]
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))[:limit]


def _date(value: Any) -> str:
    if not value: return "未记录日期"
    try:
        if isinstance(value, (int, float)): return datetime.fromtimestamp(value).strftime("%Y-%m-%d %H:%M")
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).strftime("%Y-%m-%d %H:%M")
    except (ValueError, OverflowError, OSError): return "未记录日期"


def _within(item: dict, options: dict) -> bool:
    stamp = _date(item.get("recorded_at") or item.get("date") or item.get("created_at"))
    if stamp == "未记录日期": return True  # Unknown date is disclosed, never silently excluded.
    day = stamp[:10]
    return not ((options["date_from"] and day < options["date_from"]) or (options["date_to"] and day > options["date_to"]))


def _facts(workspace: dict, profile: dict, options: dict) -> list[tuple[str, list[str]]]:
    """One common factual representation gives both formats identical content."""
    sections: list[tuple[str, list[str]]] = []
    case = workspace.get("case") or {}
    cap = 12 if options["detail"] == "brief" else 50
    if "profile" in options["sections"]:
        lines = []
        profile = dict(profile)
        if profile.get("age") is None and profile.get("birth_date"):
            try:
                birth = date.fromisoformat(profile["birth_date"])
                today = date.today()
                if birth <= today: profile["age"] = today.year - birth.year - ((today.month, today.day) < (birth.month, birth.day))
            except (ValueError, TypeError): pass
        profile["sex"] = {"male": "男", "female": "女", "other": "其他", "unknown": "未记录", "undisclosed": "未填写"}.get(profile.get("sex"), profile.get("sex"))
        for key, label in (("display_name", "姓名或昵称"), ("age", "年龄"), ("sex", "性别"), ("height_cm", "身高 cm"), ("weight_kg", "体重 kg"), ("occupation", "职业")):
            lines.append(label + "：" + _text(profile.get(key)))
        sections.append((SECTIONS["profile"], lines))
    if "questionnaire" in options["sections"]:
        q = workspace.get("questionnaire") or {}
        lines = []
        if q and _within(q, options):
            answers, triage = q.get("answers") or {}, q.get("triage") or {}
            lines += ["填写时间：" + _date(q.get("created_at")), "问卷协议：" + _text(q.get("protocol_version")),
                      "本次记录解读：" + _text(triage.get("analysis") or triage.get("message"))]
            risk_labels = {"routine": "当前回答未触发本规则中的红旗，不代表排除疾病", "review": "需要进一步人工核实", "soon": "建议尽快获得线下专业评估", "urgent": "请立即寻求急诊评估，不等待线上报告", "incomplete": "资料不足，暂不能完成安全分流"}
            risk = triage.get("risk_level") or q.get("risk_level")
            lines.append("分流状态：" + risk_labels.get(risk, _text(risk)))
            # Reuse the questionnaire's labels so reports do not expose enum codes.
            if q.get("protocol_version") == "spine-intake-v3":
                from .questionnaire import questionnaire_schema
                for section in questionnaire_schema()["sections"]:
                    if section["id"] == "start": continue
                    for field in section["fields"]:
                        fid = field["id"].split(".")
                        value = answers.get(fid[0])
                        if len(fid) == 2: value = value.get(fid[1]) if isinstance(value, dict) else None
                        if options["detail"] == "brief" and section["id"] in {"safety", "history"}: continue
                        labels = {str(v["value"]): v["label"] for v in field.get("options", [])}
                        if isinstance(value, list): value = [labels.get(str(v), v) for v in value]
                        elif value is not None: value = labels.get(str(value), value)
                        lines.append(field["label"] + "：" + _text(value))
            else:
                for key, label in (("pain_score", "自报疼痛 0 至 10 分"), ("symptom_location", "症状位置"), ("duration", "持续时间"), ("function_impacts", "日常功能影响"), ("history", "既往情况"), ("treatment", "当前处理"), ("goals", "本次目标")):
                    lines.append(label + "：" + _text(answers.get(key)))
            lines += ["问卷用于记录本人感受和安全分流，不是经过验证的诊断量表；不会计算所谓脊柱健康总分。"]
            for source in triage.get("sources", [])[:4]:
                if isinstance(source, dict): lines.append("问卷参考：" + _text(source.get("title")) + " " + _text(source.get("url")))
        else: lines.append("所选时间范围内未记录问卷。")
        sections.append((SECTIONS["questionnaire"], lines))
    if "assessments" in options["sections"]:
        lines = []
        for run in (workspace.get("agent_tool_runs") or [])[:cap]:
            if not _within(run, options): continue
            lines.append("Agent 工具任务：" + _date(run.get("created_at")) + "；状态 " + _status(run.get("state")))
            for card in (run.get("result_cards") or [])[:12]:
                lines.append(_text(card.get("title") or "工具输出") + "：" + _text(card.get("summary") or card.get("message")))
                metrics = card.get("metrics") or []
                if isinstance(metrics, dict): metrics = [{"label": k, "value": v, "unit": (card.get("metric_units") or {}).get(k, "")} for k, v in metrics.items()]
                for metric in metrics[:24]:
                    if isinstance(metric, dict): lines.append("  " + _text(metric.get("label") or metric.get("name")) + "：" + _text(metric.get("value")) + str(metric.get("unit") or ""))
                if card.get("limitations"): lines.append("限制：" + _text(card["limitations"]))
                if card.get("model_version"): lines.append("模型版本：" + _text(card["model_version"]))
                if card.get("inference_mode"):
                    mode = {"source_direct": "Pix2Pix 左半图直推", "rgb_transfer": "真实 RGB 迁移"}.get(card["inference_mode"], card["inference_mode"])
                    lines.append("推理模式：" + _text(mode))
                if card.get("review_status"): lines.append("复核状态：" + _status(card["review_status"]))
            if options["audience"] == "clinician": lines.append("任务标识：" + _text(run.get("id")))
        for entry in (workspace.get("legacy_timeline") or [])[:cap]:
            if entry.get("kind") in {"checkin", "check_in", "status", "status_checkin", "rehab", "followup", "clinical"} or not _within(entry, options): continue
            payload = entry.get("payload") or {}
            lines.append("已有检测记录：" + _date(entry.get("recorded_at") or entry.get("created_at")) + "；" + _text(entry.get("title") or entry.get("kind")))
            if entry.get("summary"): lines.append("记录摘要：" + _text(entry["summary"]))
            for key, label in (("angles", "角度测量"), ("metrics", "测量值"), ("summary", "记录摘要"), ("model_version", "模型版本"), ("quality", "质量状态")):
                if payload.get(key) is not None: lines.append(label + "：" + _text(payload[key]))
        for job in (workspace.get("gait_jobs") or [])[:cap]:
            if not _within(job, options): continue
            report = job.get("report") or job.get("result") or {}
            lines.append("步态任务：" + _date(job.get("created_at")) + "；状态 " + _status(job.get("state") or job.get("status")))
            if report.get("analysis"): lines.append("已存分析：" + _text(report["analysis"].get("summary") if isinstance(report["analysis"], dict) else report["analysis"]))
            if report.get("quality"): lines.append("质量记录：" + _text(report["quality"]))
        for item in (workspace.get("assessments") or [])[:cap]:
            if _within(item, options): lines.append("评估记录：" + _text(item.get("tool_id") or item.get("title") or item.get("modality")) + "；状态 " + _status(item.get("state")) + "；" + _text(item.get("summary") or item.get("result")))
        if not lines: lines.append("所选时间范围内未记录检测结果；报告不会补填测量数值。")
        lines.append("各模态数据单独保留。未校准的体表图、伪 DR 和模型分数不是临床诊断；协议、模型或采集条件不同的结果不得直接解释为改善或恶化。")
        sections.append((SECTIONS["assessments"], lines))
    if "followup" in options["sections"]:
        lines = []
        for entry in (workspace.get("legacy_timeline") or [])[:cap]:
            if entry.get("kind") not in {"checkin", "check_in", "status", "status_checkin", "rehab", "followup", "clinical"} or not _within(entry, options): continue
            payload = entry.get("payload") or {}
            kind_title = {"clinical": "临床复查记录", "checkin": "个人状态打卡"}.get(entry.get("kind"), entry.get("kind"))
            lines.append("随访记录：" + _date(entry.get("recorded_at") or entry.get("created_at")) + "；" + _text(entry.get("title") or kind_title))
            if entry.get("summary"): lines.append("记录摘要：" + _text(entry["summary"]))
            for key, label in (("pain_score", "自报疼痛"), ("notes", "备注"), ("symptoms", "症状"), ("cobb_angle", "记录的 Cobb 角"), ("summary", "记录摘要")):
                if payload.get(key) is not None: lines.append(label + "：" + _text(payload[key]))
        for review in (workspace.get("review_tasks") or [])[:cap]:
            if _within(review, options): lines.append("复核任务：" + _text(review.get("summary")) + "；状态 " + _status(review.get("state")))
        for plan in (workspace.get("rehab_plans") or [])[:cap]:
            if _within(plan, options): lines.append("已存计划：" + _text(plan.get("title") or plan.get("summary") or plan.get("id")) + "；状态 " + _status(plan.get("state") or plan.get("status")))
        if not lines: lines.append("所选时间范围内未记录随访、康复计划或复核任务。")
        sections.append((SECTIONS["followup"], lines))
    sections.append(("报告范围与限制", ["本报告是所选病例已记录信息的本地快照，用于资料整理和就医沟通，不是诊断证明，不替代医生面诊。未记录不等于正常。报告不会新开具药物或康复处方。", "范围：" + "、".join(SECTIONS[s] for s in options["sections"]), "记录日期范围：" + (options["date_from"] or "不限") + " 至 " + (options["date_to"] or "不限") + "；缺少日期的记录仍展示并标明。", "病例标识：" + _text(case.get("id")) + "；导出模板：" + REPORT_VERSION]))
    return sections


def _status(value: Any) -> str:
    return {"completed": "已完成", "failed": "失败", "cancelled": "已取消", "running": "运行中", "queued": "排队中", "needs_review": "待人工复核", "open": "待复核", "resolved": "已处理", "draft": "草稿", "approved": "已审核", "unreviewed": "尚未人工复核"}.get(str(value), _text(value))


def build_agent_report(workspace: dict, profile: dict | None, options: dict | None, format: str) -> bytes:
    options = validate_report_options(options)
    if format not in {"pdf", "docx"}: raise ValueError("请选择 PDF 或 DOCX 格式")
    if not isinstance(workspace, dict) or not isinstance(workspace.get("case"), dict): raise ValueError("缺少病例资料")
    _dependencies()
    sections = _facts(workspace, profile or {}, options)
    return _pdf(options, sections) if format == "pdf" else _docx(options, sections)


def _intro(options: dict) -> str:
    return ("医生沟通版" if options["audience"] == "clinician" else "个人阅读版") + " · " + ("摘要" if options["detail"] == "brief" else "详细记录") + " · 生成时间 " + datetime.now().strftime("%Y-%m-%d %H:%M")


def _pdf(options: dict, sections: list) -> bytes:
    try:
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfbase.cidfonts import UnicodeCIDFont
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    except ImportError as exc: raise ImportError("PDF 导出依赖未安装，请安装 requirements.txt") from exc
    font = "SpineAgentReportCN"
    if font not in pdfmetrics.getRegisteredFontNames():
        candidates = [os.getenv("MINDSPINE_FONT_PATH", ""), "C:/Windows/Fonts/simhei.ttf"]
        path = next((p for p in candidates if p and Path(p).is_file()), None)
        if path:
            pdfmetrics.registerFont(TTFont(font, path))
        else:
            font = "STSong-Light"
            if font not in pdfmetrics.getRegisteredFontNames(): pdfmetrics.registerFont(UnicodeCIDFont(font))
    styles = {"body": ParagraphStyle("body", fontName=font, fontSize=11, leading=17, spaceAfter=7, wordWrap="CJK"),
              "title": ParagraphStyle("title", fontName=font, fontSize=22, leading=30, spaceAfter=12, wordWrap="CJK"),
              "heading": ParagraphStyle("heading", fontName=font, fontSize=14, leading=22, spaceBefore=13, spaceAfter=8, keepWithNext=True, wordWrap="CJK")}
    p = lambda t, kind="body": Paragraph(escape(str(t)).replace("\n", "<br/>"), styles[kind])
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=(612, 792), leftMargin=54, rightMargin=54, topMargin=50, bottomMargin=48, title=options["title"], author="SpineInsight")
    story = [p(options["title"], "title"), p(_intro(options)), p("本报告只整理当前病例中已经保存的信息。请先查看问卷分流与数据限制，再阅读各项检测记录。"), Spacer(1, 4)]
    for title, lines in sections:
        story.append(p(title, "heading"))
        story.extend(p(line) for line in lines)
    def footer(canvas, document):
        canvas.saveState(); canvas.setFont(font, 8); canvas.drawString(54, 24, "SpineInsight  脊见智诊  个人健康记录"); canvas.drawRightString(558, 24, str(document.page)); canvas.restoreState()
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _docx(options: dict, sections: list) -> bytes:
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError as exc: raise ImportError("Word 导出依赖未安装，请安装 python-docx==1.2.0") from exc
    doc = Document(); section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]; style.font.name = "Microsoft YaHei"; style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_after = Pt(7)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.3
    doc.styles["Title"].font.size = Pt(22)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(13)
    doc.styles["Heading 1"].paragraph_format.keep_with_next = True
    doc.add_paragraph(options["title"], "Title")
    doc.add_paragraph(_intro(options))
    doc.add_paragraph("本报告只整理当前病例中已经保存的信息。请先查看问卷分流与数据限制，再阅读各项检测记录。")
    for title, lines in sections:
        doc.add_paragraph(title, "Heading 1")
        for line in lines: doc.add_paragraph(line)
    footer = section.footer.paragraphs[0]
    footer.add_run("SpineInsight  脊见智诊  个人健康记录    ")
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
    for run in footer.runs: run.font.size = Pt(8)
    doc.core_properties.title = options["title"]; doc.core_properties.author = "SpineInsight"
    doc.core_properties.subject = REPORT_VERSION
    buffer = io.BytesIO(); doc.save(buffer)
    return buffer.getvalue()
